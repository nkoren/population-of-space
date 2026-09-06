"""
One-off migration: converts the original "Population of Space.xlsx" spreadsheet
into the YAML source files under data/.

Usage:  python scripts/import-xlsx.py "path/to/Population of Space.xlsx"

Kept in the repo for provenance. After the migration the YAML files are the
source of truth; do NOT re-run this on top of hand-edited data without care.
"""
import sys, re, io, datetime as dt
from collections import defaultdict
import openpyxl, yaml

XLSX = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\nkore\Downloads\Population of Space.xlsx"
wb = openpyxl.load_workbook(XLSX, data_only=True)

# ---------------------------------------------------------------- nations
# Sheet ids -> ISO-3166 alpha-2 codes.
NATION_CODES = {
    1: ("RU", "Russia", "Includes the Soviet Union (1961-1991). Soviet-era cosmonauts are attributed to the modern successor state the dataset assigns them."),
    2: ("US", "United States", None), 3: ("UA", "Ukraine", None), 4: ("BY", "Belarus", None),
    5: ("CZ", "Czechia", "Includes Czechoslovakia."), 6: ("DE", "Germany", "Includes East and West Germany."),
    7: ("PL", "Poland", None), 8: ("BG", "Bulgaria", None), 9: ("HU", "Hungary", None), 10: ("VN", "Vietnam", None),
    11: ("CU", "Cuba", None), 12: ("MN", "Mongolia", None), 13: ("RO", "Romania", None), 14: ("FR", "France", None),
    15: ("IN", "India", None), 16: ("AU", "Australia", None), 17: ("CA", "Canada", None), 18: ("SA", "Saudi Arabia", None),
    19: ("NL", "Netherlands", None), 20: ("MX", "Mexico", None), 21: ("CR", "Costa Rica", None), 22: ("SY", "Syria", None),
    23: ("AZ", "Azerbaijan", None), 24: ("AF", "Afghanistan", None), 25: ("JP", "Japan", None), 26: ("GB", "United Kingdom", None),
    27: ("KZ", "Kazakhstan", None), 28: ("AT", "Austria", None), 29: ("BE", "Belgium", None), 30: ("CH", "Switzerland", None),
    31: ("IT", "Italy", None), 32: ("ES", "Spain", None), 33: ("PE", "Peru", None), 34: ("IS", "Iceland", None),
    35: ("TJ", "Tajikistan", None), 36: ("UZ", "Uzbekistan", None), 37: ("SK", "Slovakia", None), 38: ("ZA", "South Africa", None),
    39: ("IL", "Israel", None), 40: ("CN", "China", None), 41: ("BR", "Brazil", None), 42: ("SE", "Sweden", None),
    43: ("MY", "Malaysia", None), 44: ("CO", "Colombia", None), 45: ("KR", "South Korea", None), 46: ("AE", "United Arab Emirates", None),
    47: ("PT", "Portugal", None), 48: ("EG", "Egypt", None), 49: ("TR", "Turkey", None), 50: ("NO", "Norway", None),
    51: ("IR", "Iran", None), 52: ("SG", "Singapore", None),
}

DESTINATIONS = {
    1: ("free-flying", "Free-flying orbital", "Orbital flights not docked to a space station (Vostok, Gemini, Apollo Earth-orbit, most Shuttle flights, Inspiration4...)."),
    2: ("suborbital", "Suborbital", "Flights that crossed the 100 km Kármán line without reaching orbit (X-15 flights 90 and 91, Mercury-Redstone, New Shepard, SpaceShipOne). Flights that peaked below 100 km, such as Virgin Galactic, are not counted."),
    3: ("moon", "Moon", "Apollo lunar missions."),
    4: ("salyut-mir", "Salyut / Mir", "Soviet and Russian space stations: Salyut 1-7 and Mir."),
    5: ("skylab", "Skylab", "The US Skylab station (1973-74)."),
    6: ("iss", "International Space Station", "The ISS, continuously crewed since November 2000."),
    7: ("tiangong", "Tiangong", "Chinese space stations: Tiangong-1, Tiangong-2 and the Tiangong modular station."),
}

NAME_FIXES = {
    "Yuri Gargarin": "Yuri Gagarin", "Alan Shephard": "Alan Shepard", "John Glen": "John Glenn",
}

def slugify(s):
    s = re.sub(r"\(.*?\)", "", s)
    s = s.strip().lower()
    s = (s.replace("ł", "l").replace("ø", "o").replace("ß", "ss"))
    import unicodedata
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s

def parse_dob(v):
    if v is None or v == "": return None
    if isinstance(v, dt.datetime): return v.date().isoformat()
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", str(v).strip())
    if m:
        d, mo, y = map(int, m.groups()); return dt.date(y, mo, d).isoformat()
    raise ValueError(f"bad DOB {v!r}")

def parse_ids(v):
    """Crew / nationality cells: '1,23' or 1.0 or a mangled 123155185194216.0"""
    if v is None or v == "": return []
    if isinstance(v, float) and v > 999:   # Sheets coerced '123,155,185' into a number
        s = str(int(v)); assert len(s) % 3 == 0, v
        return [int(s[i:i+3]) for i in range(0, len(s), 3)]
    if isinstance(v, (int, float)): return [int(v)]
    out = []
    for x in str(v).split(","):
        x = x.strip()
        if not x: continue
        if len(x) > 3 and len(x) % 3 == 0:   # missing comma between 3-digit ids, e.g. '133145'
            out += [int(x[i:i+3]) for i in range(0, len(x), 3)]
        else:
            out.append(int(x))
    return out

def iso(d):
    return d.strftime("%Y-%m-%dT%H:%M:%SZ")

# ---------------------------------------------------------------- people
ws = wb["Astronauts"]
people, by_num = [], {}
used_ids = set()
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[1]: continue
    num = int(r[0]); name = NAME_FIXES.get(r[1].strip(), r[1].strip())
    wiki = (r[6] or "").strip() or None
    if wiki and not wiki.startswith("http"): wiki = "https://" + wiki
    pid = slugify(wiki.rsplit("/", 1)[1].replace("_", " ")) if wiki else slugify(name)
    base, n = pid, 2
    while pid in used_ids: pid = f"{base}-{n}"; n += 1
    used_ids.add(pid)
    nats = [NATION_CODES[i][0] for i in parse_ids(r[4])]
    p = {"id": pid, "name": name, "born": parse_dob(r[2]), "sex": r[3].strip(), "nationality": nats}
    if wiki: p["wiki"] = wiki
    people.append(p); by_num[num] = pid

# ---------------------------------------------------------------- flights
ws = wb["Flights"]
flights = []
used_fids = set()
LAUNCH_NATION = {1: "RU", 2: "US", 40: "CN"}
# Known corrections to the spreadsheet.
LANDING_FIXES = {  # flights that were still aloft when the sheet was last updated
    "Soyuz MS-26": ("2025-04-20T01:20:00Z", "Landing time filled in during migration from the spreadsheet placeholder; verify."),
    "SpaceX Crew 9": ("2025-03-18T21:57:00Z", "Landing time filled in during migration from the spreadsheet placeholder; verify."),
    "Shenzhou 19": ("2025-04-30T05:08:00Z", "Landing time filled in during migration from the spreadsheet placeholder; verify."),
}
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[1]: continue
    name = r[1].strip()
    fid = slugify(name); base, n = fid, 2
    while fid in used_fids: fid = f"{base}-{n}"; n += 1
    used_fids.add(fid)
    up = [by_num[i] for i in parse_ids(r[2])]
    down = [by_num[i] for i in parse_ids(r[3])]
    launch_nation = LAUNCH_NATION[int(r[7])]
    # Obvious sheet error: Soyuz 5 tagged as a US launch.
    if name.startswith(("Soyuz", "Vostok", "Voskhod")) and launch_nation != "RU":
        print("fixing launch nation for", name); launch_nation = "RU"
    landing = iso(r[5]); note = None
    if name in LANDING_FIXES: landing, note = LANDING_FIXES[name]
    wiki = (r[6] or "").strip() or None
    if wiki and not wiki.startswith("http"): wiki = "https://" + wiki
    f = {
        "id": fid, "name": name,
        "launch": iso(r[4]), "landing": landing,
        "launch_nation": launch_nation,
        "destination": DESTINATIONS[int(r[9])][0],
        "sector": "commercial" if r[10] == "private" else "government",
        "crew_up": up, "crew_down": down,
    }
    if wiki: f["wiki"] = wiki
    if note: f["note"] = note
    flights.append(f)

flights.sort(key=lambda f: f["launch"])

# Order people by first launch so appends go at the end.
first_launch = {}
for f in flights:
    for pid in f["crew_up"]:
        first_launch.setdefault(pid, f["launch"])
people.sort(key=lambda p: (first_launch.get(p["id"], "9999"), p["id"]))
never_flew = [p["id"] for p in people if p["id"] not in first_launch]
if never_flew: print("WARNING people with no flights:", never_flew)

# ---------------------------------------------------------------- write
class Dumper(yaml.SafeDumper):
    def increase_indent(self, flow=False, indentless=False):
        return super().increase_indent(flow, False)
def represent_list(dumper, data):
    flow = all(isinstance(x, (str, int, float)) or x is None for x in data)
    return dumper.represent_sequence("tag:yaml.org,2002:seq", data, flow_style=flow)
Dumper.add_representer(list, represent_list)
def represent_str(dumper, data):
    # Quote ISO timestamps/dates so YAML parsers keep them as strings.
    if re.match(r"^\d{4}-\d{2}-\d{2}", data):
        return dumper.represent_scalar("tag:yaml.org,2002:str", data, style="'")
    return dumper.represent_scalar("tag:yaml.org,2002:str", data)
Dumper.add_representer(str, represent_str)

def dump(obj, path, header):
    with io.open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(header)
        yaml.dump(obj, fh, Dumper=Dumper, allow_unicode=True, sort_keys=False, width=120)

dump([{"code": c, "name": n, **({"note": note} if note else {})} for _, (c, n, note) in sorted(NATION_CODES.items(), key=lambda kv: kv[1][1])],
     "data/nations.yaml", "# Nations referenced by people (nationality) and flights (launch_nation). ISO 3166-1 alpha-2 codes.\n")
dump([{"id": i, "name": n, "description": d} for _, (i, n, d) in DESTINATIONS.items()],
     "data/destinations.yaml", "# Destination categories for flights.\n")
dump(people, "data/people.yaml",
     "# Every person who has flown to space, ordered by first launch. See docs/DATA.md for the schema.\n")

by_year = defaultdict(list)
for f in flights: by_year[f["launch"][:4]].append(f)
import os
for y, fs in sorted(by_year.items()):
    dump(fs, f"data/flights/{y}.yaml", f"# Crewed spaceflights launched in {y}, in launch order. See docs/DATA.md for the schema.\n")

print(f"wrote {len(people)} people, {len(flights)} flights across {len(by_year)} years")
